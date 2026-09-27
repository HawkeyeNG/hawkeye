#!/usr/bin/env python3
"""Build Hawkeye-Translation-Review.xlsx: every unreviewed ha/ig/yo string, one
sheet per language, plus the Stay-for-the-count kit. Read-only on the repo."""
import collections, datetime as dt, json, os, re, subprocess, sys, unicodedata
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path('/home/elrio/hawkeye')
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else '/mnt/c/Users/HP/Downloads/Hawkeye-Translation-Review.xlsx')
TODAY = dt.date(2026, 9, 27)
CUTOFF = TODAY - dt.timedelta(days=45)
LANGS = {'ha': 'Hausa', 'ig': 'Igbo', 'yo': 'Yoruba'}
KIT = ROOT / 'docs/private/outreach/stay-for-the-count'

# ---------------------------------------------------------------- loading
def jload(p):
    return json.loads(p.read_text(encoding='utf-8'))

def blame_dates(path, repo=ROOT):
    """line number (1-based) -> date the line last changed."""
    rel = os.path.relpath(path, repo)
    out = subprocess.run(['git', '-C', str(repo), 'blame', '--line-porcelain', '--', rel],
                         capture_output=True, text=True, check=True).stdout
    dates, cur, line = {}, None, None
    for ln in out.splitlines():
        m = re.match(r'^[0-9a-f]{40} \d+ (\d+)', ln)
        if m:
            line = int(m.group(1))
        elif ln.startswith('committer-time '):
            cur = dt.date.fromtimestamp(int(ln.split()[1]))
        elif ln.startswith('\t'):
            dates[line] = cur
    return dates

KEYLINE = re.compile(r'^(\s*)("(?:[^"\\]|\\.)*")\s*:\s*(.*)$')

def keyed_lines(path, repo=ROOT, block_indent=None):
    """(block, key) -> date. block is the enclosing 2-space-indented object key
    for nested files (political.json, content-i18n.ts), else None."""
    dates = blame_dates(path, repo)
    res, block = {}, None
    for i, ln in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
        m = KEYLINE.match(ln)
        if not m:
            continue
        ind, k = len(m.group(1)), json.loads(m.group(2))
        if block_indent is not None and ind == block_indent and m.group(3).startswith('{'):
            block = k
            continue
        res[(block, k)] = dates.get(i)
    return res

entries = {L: [] for L in LANGS}   # per language: dicts

def add(L, bundle, key, en, tr, date, file):
    entries[L].append(dict(bundle=bundle, key=key, en=en, tr=tr, date=date, file=file))

skipped_same = collections.Counter()
for bundle, d, repo in [('web', ROOT / 'app/i18n', ROOT), ('native', ROOT / 'native/src/lib/i18n', ROOT),
                        ('backend', ROOT / 'backend/src/i18n', ROOT / 'backend')]:
    en = jload(d / 'en.json')
    for L in LANGS:
        p = d / f'{L}.json'
        tr, dates = jload(p), keyed_lines(p, repo)
        for k, v in tr.items():
            if k == '_meta' or not isinstance(v, str) or not v.strip() or k not in en:
                continue
            if v == en[k]:
                skipped_same[L] += 1
                continue
            add(L, bundle, k, en[k], v, dates.get((None, k)), str(p.relative_to(ROOT)))

# native content dictionary (keyed by English)
ct = ROOT / 'native/src/lib/content-i18n.ts'
src = ct.read_text(encoding='utf-8')
obj = json.loads(src[src.index('= {', src.index('export const CONTENT_I18N')) + 2: src.rindex('};') + 1])
cdates = keyed_lines(ct, ROOT, block_indent=2)
for L in LANGS:
    for enk, v in obj[L].items():
        if v == enk:
            skipped_same[L] += 1
            continue
        add(L, 'native-content', enk, enk, v, cdates.get((L, enk)), str(ct.relative_to(ROOT)))

pol = ROOT / 'app/i18n/political.json'
pobj, pdates = jload(pol), keyed_lines(pol, ROOT, block_indent=2)
for L in LANGS:
    for enk, v in pobj.get(L, {}).items():
        if v == enk:
            skipped_same[L] += 1
            continue
        add(L, 'political', enk, enk, v, pdates.get((L, enk)), str(pol.relative_to(ROOT)))

# ---------------------------------------------------------------- where it appears
TOKEN = re.compile(r'[A-Za-z0-9_-]+(?:\.[A-Za-z0-9_-]+)+')
def scan(files, keyset):
    idx = collections.defaultdict(set)
    for f in files:
        if f.stat().st_size > 1_000_000:
            continue
        txt = f.read_text(encoding='utf-8', errors='ignore')
        for t in set(TOKEN.findall(txt)) & keyset:
            idx[t].add(f)
    return idx

allkeys = {b: {e['key'] for L in LANGS for e in entries[L] if e['bundle'] == b} for b in ('web', 'native', 'backend')}
web_files = [p for p in (ROOT / 'app').rglob('*') if p.suffix in ('.html', '.js') and p.is_file()
             and not any(x in p.parts for x in ('i18n', 'vendor', 'reg', 'maps', 'press'))]
nat_files = [p for p in (ROOT / 'native/src').rglob('*') if p.suffix in ('.ts', '.tsx') and p.is_file()
             and 'i18n' not in p.parts and p.name != 'content-i18n.ts']
be_files = [p for p in (ROOT / 'backend/src').rglob('*.js') if 'i18n' not in p.parts and 'node_modules' not in p.parts]
WIDX, NIDX, BIDX = scan(web_files, allkeys['web']), scan(nat_files, allkeys['native']), scan(be_files, allkeys['backend'])

def fmt_files(fs, base):
    names = sorted({str(f.relative_to(base)) for f in fs})
    pages = [n for n in names if n.endswith('.html')]
    if len(pages) > 5:
        names = [f'{len(pages)} pages (site-wide chrome)'] + [n for n in names if not n.endswith('.html')]
    return ', '.join(names[:6]) + (f' +{len(names) - 6} more' if len(names) > 6 else '')

BACKEND_CH = {'note': 'push notification + in-app alert feed', 'tg': 'Telegram message', 'otp': 'one-time code message',
              'decl': 'result-declaration alert', 'case': 'case outcome (alert feed)', 'pday': 'Practice Day reminder (push)',
              'eday': 'election-day reminder (push)', 'pnudge': 'practice nudge (push)', 'word': 'word used inside alerts'}

# content.ts sections
ctxt = (ROOT / 'native/src/lib/content.ts').read_text(encoding='utf-8')
sec_starts = [(m.start(), m.group(1)) for m in re.finditer(r'^  ([\w-]+): \{', ctxt, re.M)]
NAT_TEXT = {f: f.read_text(encoding='utf-8', errors='ignore') for f in nat_files + [ROOT / 'native/src/lib/pages.json']
            if f.exists() and f.name != 'content.ts'}
def content_where(en):
    probe = re.split(r"['\"\\’“”]", en)[0][:60]
    if len(probe) < 12:
        probe = max(re.split(r"['\"\\’“”]", en), key=len)[:60]
    i = ctxt.find(probe)
    if i >= 0:
        sec = [s for pos, s in sec_starts if pos < i]
        sec = sec[-1] if sec else 'explainer'
        return sec, f'Native: "{sec}" page (lib/content.ts)'
    hits = sorted(str(f.relative_to(ROOT / 'native/src')) for f, t in NAT_TEXT.items() if probe in t)
    if hits:
        stem = Path(hits[0]).stem
        return stem, 'Native: ' + ', '.join(hits[:3])
    return 'explainer pages', 'Native: explainer pages'

# ---------------------------------------------------------------- area + priority
HIGH_AREAS = ['Home', 'Sign-up & sign-in', 'Report flow', 'Errors', 'Push & server messages', 'Shell & navigation']
LOW_PREFIX = {'admin', 'bench', 'press', 'tiktok', 'preview', 'train', 'train2', 'traindavina', 'trainderek', 'osun',
              'meta', 'post', 'review', 'links', 'assistant', 'ask', '404'}
WEB_AREA = {'index': 'Home', 'join': 'Sign-up & sign-in', 'lang': 'Sign-up & sign-in', 'tour': 'Sign-up & sign-in',
            'invite': 'Sign-up & sign-in', 'invite2': 'Sign-up & sign-in',
            'observe': 'Report flow', 'choose-unit': 'Report flow', 'receipt': 'Report flow', 'report': 'Report flow',
            'unit': 'Report flow', 'pu': 'Report flow', 'incidents': 'Report flow', 'case': 'Public docket & cases',
            'notifications': 'Push & server messages', 'common': 'Shell & navigation', 'nav': 'Shell & navigation',
            'time': 'Shell & navigation', 'practiceday': 'Practice day', 'practice-day': 'Practice day',
            'privacy': 'Legal: privacy & terms', 'terms': 'Legal: privacy & terms'}
NAT_AREA = [('app.tabs.index', 'Home'), ('app.welcome', 'Sign-up & sign-in'), ('app.open', 'Sign-up & sign-in'),
            ('app.sign-in', 'Sign-up & sign-in'), ('app.join', 'Sign-up & sign-in'), ('components.password-field', 'Sign-up & sign-in'),
            ('components.tour', 'Sign-up & sign-in'), ('app.report', 'Report flow'), ('app.incidents', 'Report flow'),
            ('lib.submit', 'Report flow'), ('lib.location', 'Report flow'), ('lib.register-pack', 'Report flow'),
            ('components.choose-unit', 'Report flow'), ('components.capture-camera', 'Report flow'),
            ('components.report-sheet', 'Report flow'), ('components.report-content', 'Report flow'),
            ('components.serial-field', 'Report flow'), ('components.unit-search', 'Report flow'),
            ('components.contest-picker', 'Report flow'), ('components.sheet-reference', 'Report flow'),
            ('components.wizard', 'Report flow'), ('lib.errors', 'Errors'), ('lib.api', 'Errors'),
            ('app._layout', 'Errors'), ('components.header', 'Shell & navigation'), ('common', 'Shell & navigation'),
            ('app.tabs.more', 'Shell & navigation'), ('app.tabs.alerts', 'Push & server messages'),
            ('app.assistant', '~Assistant'), ('components.ask-fab', '~Assistant'), ('components.rekor-anchor', '~Ledger anchoring')]
ERR_KEY = re.compile(r'error|fail|could-not|couldnt|cannot|cant-|invalid|try-again|went-wrong|not-found|too-many|expired|denied|timed-out|timeout')
ERR_EN = re.compile(r"\b(error|failed|failure|could ?n[o'’]t|can[’']?t|cannot|unable|invalid|try again|went wrong|not found|too many|expired|denied|timed out)\b", re.I)
PRIO_KEYS = set(jload(ROOT / 'app/i18n/ha.json')['_meta'].get('priority', []))

def pretty(s):
    return s.replace('-', ' ').replace('_', ' ').strip().capitalize()

def classify(e):
    """-> (priority, area, where)"""
    b, k = e['bundle'], e['key']
    if b == 'backend':
        pre = k.split('.')[0]
        found = fmt_files(BIDX.get(k, ()), ROOT / 'backend')
        where = 'Server: ' + BACKEND_CH.get(pre, pre) + (f' ({found})' if found else '')
        return 'High', 'Push & server messages', where
    if b == 'native-content':
        sec, where = content_where(e['en'])
        area = 'Legal: privacy & terms' if sec in ('privacy', 'terms') else f'Explainer: {sec}'
        return 'Medium', area, where
    if b == 'political':
        return 'Medium', 'Political data', 'Web: political.html · Native: Political screen (lib/political.ts)'
    pre = k.split('.')[0]
    if b == 'web':
        webarea = lambda p: WEB_AREA.get(p, ('~' if p in LOW_PREFIX else '') + pretty(p))
        area = webarea(pre)
        pages = sorted({f.stem for f in WIDX.get(k, ()) if f.suffix == '.html'})
        if pages and pre not in pages and pre not in ('common', 'nav', 'time'):
            # the key's name says one page but it only renders on others: rank by where it really shows
            cands = [webarea(p) for p in pages]
            area = min(cands, key=lambda a: (a.startswith('~'), a not in HIGH_AREAS, a))
        found = fmt_files(WIDX.get(k, ()), ROOT / 'app')
        page = ROOT / 'app' / f'{pre}.html'
        where = 'Web: ' + (found or (f'{pre}.html (by key name)' if page.exists() else f'{pre} (by key name)'))
    else:
        path = '.'.join(k.split('.')[1:-1]) if pre == 'n' else pre
        area = next((a for p, a in NAT_AREA if path == p or path.startswith(p + '.')), None)
        if area is None:
            if pre == 'n':
                seg = path.split('.')
                area = pretty(seg[1] if seg[0] in ('app', 'components', 'lib') and len(seg) > 1 else seg[0])
                if seg[:2] == ['app', 'tabs'] and len(seg) > 2:
                    area = pretty(seg[2])
            else:
                area = WEB_AREA.get(pre, ('~' if pre in LOW_PREFIX else '') + pretty(pre))
        found = fmt_files(NIDX.get(k, ()), ROOT / 'native/src')
        where = 'Native: ' + (found or (path.replace('.', '/') + ' (by key name)' if pre == 'n' else f'{pre} (by key name)'))
    low = area.startswith('~')
    area = area.lstrip('~')
    if low:
        return 'Low', area, where
    if area in HIGH_AREAS or k in PRIO_KEYS:
        return 'High', area, where
    if ERR_KEY.search(k) or ERR_EN.search(e['en']):
        return 'High', 'Errors', where
    return 'Medium', area, where

# ---------------------------------------------------------------- glossary
GLOSS = []
for ln in (ROOT / 'app/i18n/GLOSSARY.md').read_text(encoding='utf-8').splitlines():
    m = re.match(r'^\|(\d+)\|([^|]+)\|([^|]*)\|([^|]*)\|([^|]*)\|', ln)
    if m:
        GLOSS.append(dict(n=int(m.group(1)), en=m.group(2).strip(), ha=m.group(3).strip(), ig=m.group(4).strip(), yo=m.group(5).strip()))
def gloss_patterns(term):
    base = re.sub(r'\(([^)]*)\)', r'/\1', term)
    alts = [a.strip() for a in base.split('/') if a.strip() and a.strip().lower() not in ('of a result',)]
    return [re.compile(r'\b' + re.escape(a) + r's?\b', re.I) for a in alts]
for g in GLOSS:
    g['re'] = gloss_patterns(g['en'])

PH = re.compile(r'\{\{?\s*[\w.]+\s*\}?\}|<[^<>]+>')
PROTECTED = ['Hawkeye', 'IReV', 'EC8A', 'INEC']

# ---------------------------------------------------------------- group rows
PRANK = {'High': 0, 'Medium': 1, 'Low': 2}
def area_rank(a):
    return (HIGH_AREAS.index(a), '') if a in HIGH_AREAS else (99, a)

sheets = {}
for L in LANGS:
    groups = collections.OrderedDict()
    for e in entries[L]:
        if e['date'] and e['date'] < CUTOFF:
            continue
        pr, area, where = classify(e)
        g = groups.setdefault((e['en'], e['tr']), dict(en=e['en'], tr=e['tr'], members=[]))
        g['members'].append(dict(e, prio=pr, area=area, where=where))
    rows = []
    for g in groups.values():
        ms = sorted(g['members'], key=lambda m: (PRANK[m['prio']], area_rank(m['area']), m['bundle'], m['key']))
        top = ms[0]
        keys = []
        for m in ms:
            keys.append(f"{m['bundle']}: " + ('[keyed by the English text]' if m['bundle'] in ('native-content', 'political') else m['key']))
        wheres = list(dict.fromkeys(m['where'] for m in ms))
        dates = [m['date'] for m in ms if m['date']]
        notes = []
        if any(m['key'] in PRIO_KEYS for m in ms):
            notes.append("On the owner's read-first list.")
        phs = list(dict.fromkeys(PH.findall(g['en'])))
        vars_ = [p for p in phs if p.startswith('{')]
        tags = sorted({re.match(r'</?\s*(\w+)', p).group(1) for p in phs if p.startswith('<') and re.match(r'</?\s*(\w+)', p)})
        if vars_:
            notes.append('Keep exactly: ' + ' '.join(vars_))
        if tags:
            notes.append('Keep the HTML tags (' + ', '.join('<%s>' % t for t in tags) + ') and their attributes.')
        if set(PH.findall(g['tr'])) != set(PH.findall(g['en'])):
            notes.append('CHECK: placeholders/tags differ from the English.')
        prot = [p for p in PROTECTED if p in g['en']]
        missing = [p for p in prot if p not in g['tr']]
        if missing:
            notes.append('CHECK: ' + ', '.join(missing) + ' missing from the translation.')
        hits = [gl for gl in GLOSS if any(r.search(g['en']) for r in gl['re'])][:3]
        if hits:
            notes.append('Glossary: ' + '; '.join(f"{h['en']} = {h[L]}" for h in hits))
        if any(m['key'] == 'pday.weekdays' or m['key'] == 'pday.months' for m in ms):
            notes.append('Comma-separated list: keep the same number of items and the commas.')
        if top['area'] == 'Legal: privacy & terms':
            notes.append('Legal text: the GLOSSARY asks for a human translator here. Read against the English line by line.')
        rows.append(dict(prio=top['prio'], area=top['area'], where=' · '.join(wheres) if len(wheres) <= 4 else ' · '.join(wheres[:4]) + f' · +{len(wheres) - 4} more',
                         keys='\n'.join(keys), en=g['en'], tr=g['tr'], notes=notes, date=max(dates) if dates else None,
                         firstkey=top['key']))
    rows.sort(key=lambda r: (PRANK[r['prio']], area_rank(r['area']), r['firstkey']))
    # consistency: same English, different translations
    by_en = collections.defaultdict(list)
    for i, r in enumerate(rows, 1):
        r['n'] = i
        by_en[r['en']].append(i)
    for r in rows:
        others = [n for n in by_en[r['en']] if n != r['n']]
        if others:
            r['notes'].append('Same English is translated differently in row ' + ', '.join(f'#{n}' for n in others) + ' — pick one wording.')
    sheets[L] = rows

# ---------------------------------------------------------------- kit
def paras(p):
    return [x.strip() for x in re.split(r'\n\s*\n', p.read_text(encoding='utf-8').strip()) if x.strip()]
kit_rows = []
wa = {L: paras(KIT / f'whatsapp-{L}.txt') for L in ['en', *LANGS]}
for i in range(max(len(v) for v in wa.values())):
    kit_rows.append((f'WhatsApp message ¶{i + 1}', {L: (wa[L][i] if i < len(wa[L]) else '') for L in wa}, ''))
node = subprocess.run(['node', '-e', r'''
const fs=require('fs');const s=fs.readFileSync(process.argv[1],'utf8');
const a=s.indexOf('const CONTENT = ')+'const CONTENT = '.length;const e=s.indexOf('\n};',a)+2;
const C=Function('return ('+s.slice(a,e)+')')();
const flat=(o,p,out)=>{if(typeof o==='string')out.push([p,o]);else if(Array.isArray(o))o.forEach((v,i)=>flat(v,p+'['+(i+1)+']',out));else if(o&&typeof o==='object')for(const[k,v]of Object.entries(o))flat(v,p?p+'.'+k:k,out);return out;};
const r={};for(const[l,o]of Object.entries(C))r[l]=flat(o,'',[]).filter(([p])=>p!=='name');
console.log(JSON.stringify(r));''', str(KIT / 'build.mjs')], capture_output=True, text=True, check=True)
fly = json.loads(node.stdout)
fmap = {L: dict(fly[L]) for L in fly}
for path, en in fly['en']:
    note = 'Braces mark the gold highlight: keep one pair around the key word.' if '{' in en else ''
    if path == 'title':
        note = 'Browser/PDF title only.'
    kit_rows.append((f'Flyer: {path}', {L: fmap[L].get(path, '') for L in fmap}, note))
def split_lists(ps):
    out = []
    for p in ps:
        lines = p.splitlines()
        out += lines if len(lines) > 1 and all(l.lstrip().startswith('- ') for l in lines) else [p]
    return out
talk = {L: paras(KIT / f'talk-{L}.md') for L in ['en', *LANGS]}
split = {L: split_lists(v) for L, v in talk.items()}
if len({len(v) for v in split.values()}) == 1:     # bullets line up in every language: one row per bullet
    talk = split
else:
    print('WARNING: bullet counts differ between languages; talk lists kept as blocks')
for i in range(max(len(v) for v in talk.values())):
    en = talk['en'][i] if i < len(talk['en']) else ''
    note = 'Safety part: the talk says to keep it word for word — accuracy matters most here.' if re.search(r'safe|No result is worth', en, re.I) else ''
    kit_rows.append((f'Talk ¶{i + 1}', {L: (talk[L][i] if i < len(talk[L]) else '') for L in talk}, note))

# ---------------------------------------------------------------- write
GREEN, WHITE = '004225', 'FFFFFF'
HFONT, HFILL = Font(bold=True, color=WHITE), PatternFill('solid', fgColor=GREEN)
PFILL = {'High': PatternFill('solid', fgColor='F8D7DA'), 'Medium': PatternFill('solid', fgColor='FFF3CD'), 'Low': PatternFill('solid', fgColor='E9ECEF')}
WRAP = Alignment(wrap_text=True, vertical='top')
ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

def put(ws, r, c, v):
    if isinstance(v, str) and ILLEGAL.search(v):
        raise SystemExit(f'illegal control character in cell {ws.title}!{r},{c}: {v[:60]!r}')
    cell = ws.cell(row=r, column=c, value=v)
    if isinstance(v, str):
        cell.data_type = 's'          # never let a leading "=" become a formula
    cell.alignment = WRAP
    return cell

def header(ws, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        cell = put(ws, 1, c, h)
        cell.font, cell.fill = HFONT, HFILL
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = 'A2'

wb = Workbook()
how = wb.active
how.title = 'How to review'
COLS = ['#', 'Priority', 'Where it appears', 'Key', 'English', 'Current translation', 'Reviewer: OK / Fix', 'Corrected text', 'Notes', 'Last changed']
WID = [6, 9, 30, 34, 55, 55, 12, 55, 42, 12]
for L, name in LANGS.items():
    ws = wb.create_sheet(name)
    header(ws, COLS, WID)
    for i, r in enumerate(sheets[L], 2):
        vals = [r['n'], r['prio'], r['where'], r['keys'], r['en'], r['tr'], None, None, '\n'.join(r['notes']) or None,
                r['date'].isoformat() if r['date'] else None]
        for c, v in enumerate(vals, 1):
            put(ws, i, c, v)
        ws.cell(row=i, column=2).fill = PFILL[r['prio']]
        ws.cell(row=i, column=7).fill = PatternFill('solid', fgColor='E8F4EA')
        ws.cell(row=i, column=8).fill = PatternFill('solid', fgColor='E8F4EA')
    last = len(sheets[L]) + 1
    ws.auto_filter.ref = f'A1:J{last}'
    dv = DataValidation(type='list', formula1='"OK,Fix"', allow_blank=True, showErrorMessage=True,
                        errorTitle='OK or Fix', error='Choose OK or Fix from the list.')
    ws.add_data_validation(dv)
    dv.add(f'G2:G{last}')

kit = wb.create_sheet('Kit')
KCOLS = ['#', 'Source', 'English']
KWID = [5, 22, 48]
for name in LANGS.values():
    KCOLS += [name, f'{name}: OK / Fix', f'{name}: corrected']
    KWID += [48, 11, 48]
KCOLS.append('Notes')
KWID.append(34)
header(kit, KCOLS, KWID)
kit.freeze_panes = 'D2'
kdv = DataValidation(type='list', formula1='"OK,Fix"', allow_blank=True)
kit.add_data_validation(kdv)
for i, (src_, t, note) in enumerate(kit_rows, 2):
    vals = [i - 1, src_, t['en']]
    for L in LANGS:
        vals += [t.get(L, ''), None, None]
    vals.append(note or None)
    for c, v in enumerate(vals, 1):
        put(kit, i, c, v)
klast = len(kit_rows) + 1
kit.auto_filter.ref = f'A1:{kit.cell(row=1, column=len(KCOLS)).column_letter}{klast}'
for j, _ in enumerate(LANGS):
    col = kit.cell(row=1, column=5 + 3 * j).column_letter
    kdv.add(f'{col}2:{col}{klast}')

# How to review
cnt = {L: collections.Counter(r['prio'] for r in sheets[L]) for L in LANGS}
lines = [
    ('Hawkeye translation review', 'title'),
    ('Thank you for reviewing. Hawkeye is a free, independent app that lets observers photograph the result posted at their polling unit. '
     'Most Hausa, Igbo and Yorùbá text in it was drafted by an AI model and no native speaker has read it yet. Your job is to make it read '
     'naturally to an ordinary person in your language.', None),
    ('What to do', 'h'),
    ('1. Open the sheet for your language (Hausa, Igbo or Yoruba). Start at the top: rows are sorted High → Medium → Low, then by screen.', None),
    ('2. Read "English" and "Current translation". If the translation is right and natural, choose OK in "Reviewer: OK / Fix".', None),
    ('3. If it needs changing, choose Fix and type the WHOLE corrected sentence in "Corrected text" (not just the changed word).', None),
    ('4. Use "Notes" for questions or doubts. Please do not edit the Key, English or Current translation columns — they are how your corrections find their way back into the app.', None),
    ('5. One row can cover several places in the app (several keys). Your correction is applied to all of them.', None),
    ('6. The "Kit" sheet holds the Stay-for-the-count WhatsApp message, flyer and talk. Review only your language\'s columns there.', None),
    ('7. Save as .xlsx and send the file back. Partial reviews are welcome: High rows first matters most.', None),
    ('Rules', 'h'),
    ('• Keep placeholders exactly as they are: {v0}, {v1}, {name}, {date}, {code} and the like. They are replaced by numbers and names when the app runs. '
     'You may move them within the sentence, but never translate, respell or delete them. Corrections that change or drop one are rejected automatically.', None),
    ('• Keep HTML tags such as <strong>…</strong> or <a href="…">…</a> and their attributes; translate only the words between them.', None),
    ('• Keep product and institution names as written: Hawkeye, IReV, EC8A, INEC (also Telegram, WhatsApp, Google Play, App Store).', None),
    ('• Numbers stay in digits. Never write numbers out as words.', None),
    ('• Aim for plain, natural wording an ordinary voter would use — not a word-for-word copy of the English. Keep it short where the English is short (buttons, labels).', None),
    ('• Use the same term for the same thing everywhere. The glossary below is the agreed vocabulary; "Notes" tells you when a row uses a glossary term, '
     'and flags rows where the same English was translated two different ways.', None),
    ('• "Accredited voters" and "Registered voters" are different numbers on the same sheet. They must never share one word.', None),
    ('• "Unverified" carries legal weight: it must be unmistakable, not softened.', None),
    ('• Privacy and terms text is legal wording. Read it line by line against the English.', None),
    ('Priority', 'h'),
    ('High: home, sign-up/sign-in, the report flow, error messages, and push/Telegram/alert texts the server sends. Medium: most other screens. '
     'Low: admin and rarely-seen screens.', None),
    ('Glossary (app/i18n/GLOSSARY.md) — provisionally checked by one speaker per language, 2026-09-07', 'h'),
]
import math
HOW_W = [5, 30, 34, 34, 34]
for col, w in zip('ABCDE', HOW_W):
    how.column_dimensions[col].width = w
def para(r, text, kind):
    c = put(how, r, 1, text)
    how.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    if kind == 'title':
        c.font = Font(bold=True, size=16, color=GREEN)
        how.row_dimensions[r].height = 26
    elif kind == 'h':
        c.font = Font(bold=True, size=12, color=GREEN)
        how.row_dimensions[r].height = 22
    else:
        how.row_dimensions[r].height = 15.5 * max(1, math.ceil(len(text) / 125)) + 3
r = 1
for text, kind in lines:
    para(r, text, kind)
    r += 1
for c, h in enumerate(['#', 'English', 'Hausa', 'Igbo', 'Yorùbá'], 1):
    cell = put(how, r, c, h)
    cell.font, cell.fill = HFONT, HFILL
r += 1
for g in GLOSS:
    for c, v in enumerate([g['n'], g['en'], g['ha'], g['ig'], g['yo']], 1):
        put(how, r, c, v)
    r += 1
tail = [
    ('Yorùbá reviewers: glossary rows 13 (Gbe jade), 15 (Iṣẹlẹ), 16 (Rira ibo) and 25 (Nomba siriali) have no tone marks while the rest do. '
     'If that was accidental, say so in Notes — the app follows the glossary.', None),
    ('Scope of this review', 'h'),
    (f'No human review of the interface strings is recorded (only the 25-term glossary above, provisional, 2026-09-07; every bundle is marked '
     f'"machine-draft"). So every ha/ig/yo value that changed in the 45 days before {TODAY.isoformat()} counts as unreviewed: cutoff {CUTOFF.isoformat()}. '
     f'All translations were first added on or after 2026-09-08, so every translated string is included. Strings identical to English '
     f'(URLs, handles, IDs, placeholder-only strings) are left out.', None),
    ('Sources: web app/i18n/*.json and political.json; native native/src/lib/i18n/*.json and lib/content-i18n.ts; server backend/src/i18n/*.json; '
     'kit docs/private/outreach/stay-for-the-count/. "Last changed" is the git date of that line.', None),
    ('Rows per sheet: ' + '; '.join(f"{LANGS[L]} {len(sheets[L])} (High {cnt[L]['High']}, Medium {cnt[L]['Medium']}, Low {cnt[L]['Low']})" for L in LANGS)
     + f'; Kit {len(kit_rows)}.', None),
]
for text, kind in tail:
    para(r, text, kind)
    r += 1
wb.save(OUT)

print('cutoff', CUTOFF, 'skipped identical-to-English', dict(skipped_same))
for L in LANGS:
    print(LANGS[L], len(sheets[L]), dict(cnt[L]), 'entries', len(entries[L]),
          'undated', sum(1 for e in entries[L] if not e['date']),
          'oldest', min(e['date'] for e in entries[L] if e['date']))
    print('  areas', collections.Counter((r['prio'], r['area']) for r in sheets[L]).most_common(60))
print('kit rows', len(kit_rows), 'saved', OUT)
