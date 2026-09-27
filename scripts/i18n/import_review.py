#!/usr/bin/env python3
"""Apply a returned translation-review spreadsheet back to the bundles.

    python3 scripts/i18n/import_review.py Hawkeye-Translation-Review.xlsx            # dry run
    python3 scripts/i18n/import_review.py Hawkeye-Translation-Review.xlsx --apply    # write

Reads the Hausa / Igbo / Yoruba sheets. A row is applied when "Reviewer: OK / Fix"
is Fix and "Corrected text" is filled; the correction goes to EVERY key listed in
the row's Key cell ("web: key", "native: key", "backend: key", or
"native-content:" / "political:" for the dictionaries keyed by the English text).

Targets:
    web            app/i18n/<lang>.json
    native         native/src/lib/i18n/<lang>.json
    backend        backend/src/i18n/<lang>.json
    native-content native/src/lib/content-i18n.ts  (+ its builder input tmp/content_tr_*.json)
    political      app/i18n/political.json         (+ its builder input tmp/political_tr.json)

content-i18n.ts and political.json are GENERATED from gitignored tmp/ files; both
are patched so a rebuild does not silently revert a reviewer's fix.

Guards (a refused row changes nothing):
  * a placeholder ({v0}, {name}, {{count}}) or HTML tag added, changed or dropped
    -- the correction's set must match the English or the current translation;
  * Hawkeye / IReV / EC8A / INEC dropped where the current text has it;
  * the bundle no longer holds the "Current translation" the sheet was built from;
  * two rows giving the same key different corrections.
Values are NFC-normalised. Files are edited in place, value by value, so key order,
indentation and escaping stay exactly as they were; every edited file is re-parsed
and compared against the expected result before anything is written.

Exit status: 0 when nothing was refused, 1 otherwise (dry run or --apply).
The Kit sheet is prose in docs/; its Fix rows are listed for applying by hand.
"""
import argparse
import collections
import json
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl is required: python3 -m pip install openpyxl (or use a venv)')

SHEETS = {'Hausa': 'ha', 'Igbo': 'ig', 'Yoruba': 'yo', 'Yorùbá': 'yo'}
PH = re.compile(r'\{\{?\s*[\w.]+\s*\}?\}|<[^<>]+>')
PROTECTED = ('Hawkeye', 'IReV', 'EC8A', 'INEC')
LIST_KEYS = {'pday.weekdays', 'pday.months'}          # comma-separated lists in the backend bundle
KEY_LINE = re.compile(r'^\s*(web|native|backend|native-content|political)\s*:\s*(.*?)\s*$')


def nfc(s):
    return unicodedata.normalize('NFC', s)


def enc(s):
    return json.dumps(s, ensure_ascii=False)


# ---------------------------------------------------------------- files
class Doc:
    """One bundle file, edited as text so its formatting survives."""

    def __init__(self, path, kind):
        self.path, self.kind = path, kind            # kind: 'json' | 'ts'
        with open(path, encoding='utf-8', newline='') as f:
            self.orig = f.read()
        self.text = self.orig
        self.data = self.parse(self.orig)
        self.expected = json.loads(json.dumps(self.data))
        self.edits = 0

    def parse(self, text):
        if self.kind == 'json':
            return json.loads(text)
        head = text.index('= {', text.index('export const CONTENT_I18N')) + 2
        return json.loads(text[head:text.rindex('};') + 1])

    def get(self, block, key):
        node = self.data[block] if block is not None else self.data
        return node.get(key) if isinstance(node, dict) else None

    def set(self, block, key, old, new):
        """Replace `"key": "old"` (inside the 2-space-indented `"block": {` when given)."""
        start, end, indent = 0, len(self.text), '  '
        if block is not None:
            opener = '\n  ' + enc(block) + ': {'
            start = self.text.find(opener)
            if start < 0 or self.text.find(opener, start + 1) >= 0:
                raise ValueError(f'block {block[:40]!r} not found exactly once')
            end = self.text.find('\n  }', start) + 1       # keep the newline: the block's last entry ends there
            indent = '    '
        for encode in (enc, lambda s: json.dumps(s)):
            needle = '\n' + indent + encode(key) + ': ' + encode(old)
            hits = [m.start() for m in re.finditer(re.escape(needle) + r'(?=,?\r?\n)', self.text[start:end])]
            if len(hits) == 1:
                i = start + hits[0]
                repl = '\n' + indent + encode(key) + ': ' + encode(new)
                self.text = self.text[:i] + repl + self.text[i + len(needle):]
                node = self.expected[block] if block is not None else self.expected
                node[key] = new
                self.edits += 1
                return
        raise ValueError(f'{key[:50]!r} with its current value not found exactly once')

    def verify(self):
        got = self.parse(self.text)
        if got != self.expected:
            raise ValueError('re-parsed file differs from the expected result')
        a, b = self.orig.splitlines(), self.text.splitlines()
        changed = sum(x != y for x, y in zip(a, b))
        if len(a) != len(b) or changed != self.edits:
            raise ValueError(f'formatting drift: {changed} lines changed for {self.edits} edits')

    def write(self):
        with open(self.path, 'w', encoding='utf-8', newline='') as f:
            f.write(self.text)


class Bundles:
    def __init__(self, root):
        self.root, self.docs = root, {}

    def doc(self, rel, kind='json'):
        if rel not in self.docs:
            p = self.root / rel
            self.docs[rel] = Doc(p, kind) if p.exists() else None
        return self.docs[rel]

    def targets(self, bundle, key, lang):
        """[(Doc, block, key)] -- the first is the live file, any others are builder mirrors."""
        if bundle in ('web', 'native', 'backend'):
            rel = {'web': 'app/i18n', 'native': 'native/src/lib/i18n', 'backend': 'backend/src/i18n'}[bundle]
            return [(self.doc(f'{rel}/{lang}.json'), None, key)]
        if bundle == 'native-content':
            out = [(self.doc('native/src/lib/content-i18n.ts', 'ts'), lang, key)]
            for p in sorted((self.root / 'tmp').glob('content_tr_*.json')):
                d = self.doc(str(p.relative_to(self.root)))
                if d and key in d.data:
                    out.append((d, key, lang))
            return out
        out = [(self.doc('app/i18n/political.json'), lang, key)]
        d = self.doc('tmp/political_tr.json')
        if d and key in d.data:
            out.append((d, key, lang))
        return out


# ---------------------------------------------------------------- guards
def guard(english, current, corrected, keys):
    bag = lambda s: collections.Counter(PH.findall(s))
    got = bag(corrected)
    if got != bag(english) and got != bag(current):
        want = bag(current) if bag(current) else bag(english)
        lost, extra = sorted((want - got).elements()), sorted((got - want).elements())
        return 'placeholder changed: ' + '; '.join(
            x for x in [('missing ' + ' '.join(lost)) if lost else '', ('unexpected ' + ' '.join(extra)) if extra else ''] if x)
    for name in PROTECTED:
        if name in english and name in current and name not in corrected:
            return f'"{name}" was removed'
    if any(k in LIST_KEYS for _, k in keys) and corrected.count(',') != current.count(','):
        return 'list item count changed (commas)'
    return None


# ---------------------------------------------------------------- sheet
def header_map(ws):
    cols = {}
    for i, v in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))):
        h = str(v or '').strip().lower()
        for name, pre in (('n', '#'), ('key', 'key'), ('en', 'english'), ('cur', 'current translation'),
                          ('rev', 'reviewer'), ('fix', 'corrected text')):
            if h.startswith(pre) and name not in cols:
                cols[name] = i
    missing = {'key', 'en', 'cur', 'rev', 'fix'} - cols.keys()
    if missing:
        raise SystemExit(f'sheet {ws.title!r}: missing column(s) {sorted(missing)}')
    return cols


def cell(row, i):
    v = row[i] if i is not None and i < len(row) else None
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace('\r\n', '\n')


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('xlsx')
    ap.add_argument('--apply', action='store_true', help='write the changes (default: dry run)')
    ap.add_argument('--root', default=str(Path(__file__).resolve().parents[2]), help='repo root (default: this repo)')
    ap.add_argument('--lang', default='ha,ig,yo')
    a = ap.parse_args()
    root, langs = Path(a.root), set(a.lang.split(','))
    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    B = Bundles(root)

    planned = {}                     # (docpath, block, key) -> (new, label)
    refused, notes, changes = [], [], []
    for sheet, lang in SHEETS.items():
        if sheet not in wb.sheetnames or lang not in langs:
            continue
        ws = wb[sheet]
        cols = header_map(ws)
        for rno, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if cell(row, cols['rev']).strip().lower() != 'fix':
                continue
            label = f"{sheet} row {cell(row, cols.get('n')) or rno}"
            english, current, raw = cell(row, cols['en']), cell(row, cols['cur']), cell(row, cols['fix'])
            if not raw.strip():
                notes.append(f'{label}: marked Fix but "Corrected text" is empty -- skipped')
                continue
            lead = current[:len(current) - len(current.lstrip())]
            trail = current[len(current.rstrip()):]
            corrected = nfc(lead + raw.strip() + trail)
            keys = []
            for ln in cell(row, cols['key']).splitlines():
                m = KEY_LINE.match(ln)
                if m:
                    b = m.group(1)
                    keys.append((b, english if b in ('native-content', 'political') else m.group(2)))
            if not keys:
                refused.append(f'{label}: no readable key in the Key column')
                continue
            if corrected == nfc(current):
                notes.append(f'{label}: correction equals the current text -- nothing to do')
                continue
            why = guard(english, current, corrected, keys)
            if why:
                refused.append(f'{label}: REFUSED, {why}\n      corrected: {corrected[:160]}')
                continue
            row_plan, bad = [], None
            for bundle, key in keys:
                try:
                    tg = B.targets(bundle, key, lang)
                except Exception as e:                       # noqa: BLE001
                    bad = f'{bundle}:{key}: {e}'
                    break
                if not tg or tg[0][0] is None:
                    bad = f'{bundle}:{key[:50]}: bundle file not found under {root}'
                    break
                live, block, k = tg[0]
                have = live.get(block, k)
                if have is None or nfc(have) != nfc(current):
                    bad = f'{bundle}:{key[:50]}: bundle no longer holds the sheet\'s "Current translation" (changed since export?)'
                    break
                for d, blk, kk in tg:
                    if d is not None and d.get(blk, kk) is not None:
                        row_plan.append((d, blk, kk))
            if bad:
                refused.append(f'{label}: REFUSED, {bad}')
                continue
            for d, blk, kk in row_plan:
                planned.setdefault((str(d.path), blk, kk), []).append((corrected, label, d))
            changes.append((label, keys, current, corrected))

    # two rows writing different text to one value: refuse both
    clash = set()
    for (path, blk, kk), vals in planned.items():
        if len({v[0] for v in vals}) > 1:
            labels = sorted({v[1] for v in vals})
            clash.update(labels)
            refused.append(f'{" / ".join(labels)}: REFUSED, conflicting corrections for {kk[:50]} in {path}')
    changes = [c for c in changes if c[0] not in clash]

    # apply in memory, verify every touched file
    errors = []
    for (path, blk, kk), vals in planned.items():
        new, label, d = vals[0]
        if any(v[1] in clash for v in vals):
            continue
        try:
            d.set(blk, kk, d.get(blk, kk), new)
        except ValueError as e:
            errors.append(f'{label}: {path}: {e}')
    touched = [d for d in B.docs.values() if d is not None and d.edits]
    for d in touched:
        try:
            d.verify()
        except ValueError as e:
            errors.append(f'{d.path}: {e}')

    # Kit: prose, applied by hand
    kit = []
    if 'Kit' in wb.sheetnames:
        ws = wb['Kit']
        hdr = [str(v or '') for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            for i, h in enumerate(hdr):
                if h.endswith('OK / Fix') and str(row[i] or '').strip().lower() == 'fix':
                    kit.append(f'{row[1]} [{h.split(":")[0]}]: {str(row[i + 1] or "")[:100]}')

    for label, keys, cur, new in changes:
        print(f'{label}: ' + ', '.join(f'{b}:{k[:48]}' for b, k in keys))
        print(f'   - {cur}\n   + {new}')
    for d in touched:
        print(f'  {d.edits:3d} value(s)  {d.path.relative_to(root)}')
    for n in notes:
        print('note:', n)
    for r in refused + errors:
        print(r)
    if kit:
        print(f'Kit: {len(kit)} Fix cell(s) to apply by hand in docs/private/outreach/stay-for-the-count/:')
        for k in kit:
            print('   ', k)
    ok = not refused and not errors
    if a.apply and not errors:
        for d in touched:
            d.write()
        print(f'APPLIED: {len(changes)} correction(s) -> {sum(d.edits for d in touched)} value(s) in {len(touched)} file(s).')
    elif a.apply:
        print('NOTHING WRITTEN: fix the errors above first.')
    else:
        print(f'DRY RUN: {len(changes)} correction(s) would change {sum(d.edits for d in touched)} value(s) in '
              f'{len(touched)} file(s). Re-run with --apply to write.')
    print(f'{len(refused)} refused, {len(errors)} error(s).')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
